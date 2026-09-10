import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_proposal_docx(output_path):
    doc = Document()

    # Set Margins (1 inch everywhere)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Configure Heading 1 Style
    h1_style = doc.styles['Heading 1']
    h1_style.font.name = 'Arial'
    h1_style.font.size = Pt(12)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    h1_style.paragraph_format.space_before = Pt(12)
    h1_style.paragraph_format.space_after = Pt(4)
    h1_style.paragraph_format.keep_with_next = True

    # Primary Title Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("PROJECT PROPOSAL FORM")
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy
    title_p.paragraph_format.space_after = Pt(2)

    sub_title = doc.add_paragraph()
    sub_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_title.add_run("Faculty of Engineering | Robotics and Automation Engineering")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    sub_title.paragraph_format.space_after = Pt(16)

    def add_section_header(number_str, title_str):
        h = doc.add_paragraph(style='Heading 1')
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        
        run_num = h.add_run(f"{number_str}. ")
        run_num.font.bold = True
        run_num.font.size = Pt(12)
        run_num.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
        run_title = h.add_run(title_str)
        run_title.font.bold = True
        run_title.font.size = Pt(12)
        run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return h

    # 1. Project Title
    add_section_header("1", "Project Title (Subject to Advisor Approval)")
    
    t_en = doc.add_paragraph()
    t_en.paragraph_format.left_indent = Inches(0.25)
    r = t_en.add_run("English Title: ")
    r.font.bold = True
    t_en.add_run("Cloud-Based Over-The-Air (OTA) Firmware Management System for Industrial Robot Fleet")

    t_th = doc.add_paragraph()
    t_th.paragraph_format.left_indent = Inches(0.25)
    r = t_th.add_run("Thai Title (Translated): ")
    r.font.bold = True
    t_th.add_run("ระบบจัดการและอัปเดตเฟิร์มแวร์แบบโอทีเอผ่านคลาวด์สำหรับฝูงหุ่นยนต์อุตสาหกรรม")

    # 2. Background and Rationale
    add_section_header("2", "Background and Rationale")
    bg_p1 = doc.add_paragraph()
    bg_p1.add_run(
        "In modern Industry 4.0 manufacturing environments, heterogeneous industrial robot fleets—including articulated arms, "
        "SCARA, Cartesian, and autonomous mobile robots (AGVs)—operate continuously within mission-critical production lines. "
        "Maintaining these robotic systems requires frequent firmware updates to calibrate motion-control kinematics, "
        "patch security vulnerabilities, and optimize operational routines. However, conventional industrial firmware "
        "deployment paradigms face two fundamental engineering bottlenecks:"
    )

    bottlenecks = [
        ("The Blast Radius and Production Halt Risk: ", 
         "Traditional global \"all-at-once\" deployments risk propagating unvetted firmware bugs to an entire fleet simultaneously, causing catastrophic plant shutdowns (\"bricking\" the fleet) and massive operational losses."),
        ("Cyber-Physical Security Vulnerabilities: ", 
         "Remote over-the-air communication channels expose industrial controllers to malicious binary injection, bit-flip corruption, and untrusted execution if end-to-end cryptographic verification is absent at the edge.")
    ]
    for bold_text, normal_text in bottlenecks:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.left_indent = Inches(0.4)
        bp.paragraph_format.space_after = Pt(3)
        r_b = bp.add_run(bold_text)
        r_b.font.bold = True
        bp.add_run(normal_text)

    bg_p2 = doc.add_paragraph()
    bg_p2.add_run(
        "To resolve these challenges, this research proposes a robust, fault-tolerant OTA firmware management architecture. "
        "The platform introduces a deterministic Canary Phased Rollout Engine (20% → 60% → 100%) designed to restrict failure "
        "exposure to a minor sentinel group with automated rollback capabilities. This is integrated with a lightweight "
        "ECDSA NIST P-256 asymmetric cryptographic verification pipeline executed directly on edge robot controllers prior to flashing. "
        "This approach transforms industrial firmware delivery from a high-risk manual intervention into an autonomous, fault-tolerant, "
        "and zero-trust engineering operation."
    )

    # 3. Objectives
    add_section_header("3", "Objectives")
    objs = [
        "To architect and develop a centralized, cloud-managed OTA orchestration framework designed for continuous, zero-downtime deployment across heterogeneous industrial robot fleets.",
        "To engineer and evaluate an edge-level cryptographic code-verification pipeline using SHA-256 and ECDSA NIST P-256, assessing its efficacy in rejecting unauthorized or tampered binaries and measuring its computational verification overhead on industrial edge controllers.",
        "To design and evaluate a Canary Phased Rollout engine governed by a deterministic Finite State Machine (FSM) with automated failure isolation and rollback mechanisms to preserve fleet operational continuity during firmware fault scenarios."
    ]
    for idx, obj in enumerate(objs, 1):
        op = doc.add_paragraph()
        op.paragraph_format.left_indent = Inches(0.25)
        op.paragraph_format.space_after = Pt(3)
        r_num = op.add_run(f"3.{idx} ")
        r_num.font.bold = True
        op.add_run(obj)

    # 4. Scope of the Study
    add_section_header("4", "Scope of the Study")
    scopes = [
        ("Kinematic Fleet Representation: ", "Five distinct industrial robotic classes: SCARA (scara-v1), Delta (delta-v2), 6-Axis Articulated (articulated-v3), Cartesian (cartesian-v1), and Automated Guided Vehicles - AGV (agv-v1) distributed across simulated factory nodes."),
        ("Edge Robot Architecture: ", "Formal Finite State Machine (FSM) managing transition cycles: Idle → Downloading → Verifying → Installing → Rebooting → Online / Rollback."),
        ("Canary Orchestration Engine: ", "Algorithmic division of fleet populations into progressive tranches (20% → 60% → 100%) with automated circuit-breaking triggers upon exceeding predefined error thresholds."),
        ("Cloud and Storage Infrastructure: ", "Go Fiber v3 backend control plane, PostgreSQL 16 and MinIO S3-compatible Object Storage for fleet metadata, deployment audit trails, and firmware binary delivery via time-limited Presigned URLs, alongside an EMQX MQTT 5.0 (QoS 1) broker for low-latency command orchestration."),
        ("Auxiliary Data Processing & Machine Learning: ", "Post-experiment exploratory analysis using Scikit-learn (Random Forest Regression for latency feature importance and Isolation Forest for outlier detection) evaluated on a synthetic simulation dataset (N = 290 trials) as an analytical study tool rather than an active real-time runtime component."),
        ("Simulation-Based Evaluation Matrix: ", "Systematic evaluation across 5 operational scenarios (N = 290 repeated simulation trials) covering nominal baseline execution, cryptographic threat mitigation, fault resilience A/B benchmarking, network stress conditions (150-500 ms latency, 0-10% packet loss), and cross-hardware compatibility verification.")
    ]
    for bold_text, normal_text in scopes:
        sp = doc.add_paragraph(style='List Bullet')
        sp.paragraph_format.left_indent = Inches(0.4)
        sp.paragraph_format.space_after = Pt(3)
        r_b = sp.add_run(bold_text)
        r_b.font.bold = True
        sp.add_run(normal_text)

    # 5. Relevant Courses
    add_section_header("5", "Relevant Courses (Multiple Selections)")
    
    # ML is now UNCHECKED
    course_list = [
        ("[  ] Manufacturing Cost Analysis and Budgeting", "[X] Introduction to Robotics and Automation"),
        ("[X] Introduction to Probability and Statistics",  "[  ] Robot Structure Design"),
        ("[  ] Computer-aided Design and Manufacturing",    "[X] Industrial Sensors and Actuators"),
        ("[  ] Programmable Logic Controller & Automation",  "[X] Software Development for Robotics & Automation"),
        ("[X] Embedded System Design",                      "[  ] Electronic Circuits"),
        ("[  ] Manufacturing Processes",                    "[  ] Principles of Digital Circuit"),
        ("[  ] Fundamentals of Mechatronics",               "[X] Basic Control Theory"),
        ("[  ] Fundamentals of Machine Learning",           "[X] Industrial Robot"),
        ("[  ] Hydraulics and Pneumatics",                  "[X] Other: Distributed Systems, Cryptography & Cloud Computing")
    ]
    
    tbl = doc.add_table(rows=len(course_list), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, (c1, c2) in enumerate(course_list):
        row = tbl.rows[row_idx]
        cell_1, cell_2 = row.cells[0], row.cells[1]
        cell_1.text = c1
        cell_2.text = c2
        for cell in (cell_1, cell_2):
            cell.width = Inches(3.25)
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if "[X]" in cell.text:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 6. Prerequisite Knowledge and Core Technologies
    doc.add_page_break()
    add_section_header("6", "Prerequisite Knowledge and Core Technologies")
    techs = [
        ("Cyber-Physical Security & Cryptography: ", "Asymmetric elliptic-curve cryptography (ECDSA NIST P-256 over secp256r1), cryptographic hash functions (SHA-256), and Public-Key Infrastructure (PKI)."),
        ("Distributed Systems & Industrial Communication: ", "Industrial IoT protocol standards (MQTT 5.0 with QoS 1 over TLS 1.2+), deterministic Finite State Machine (FSM) modeling, RESTful control plane architecture, and S3-compatible binary distribution via Presigned URLs."),
        ("Simulation Benchmarking & Statistical Evaluation: ", "Non-parametric hypothesis testing (Mann-Whitney U, Shapiro-Wilk normality tests), confidence interval estimation (95% CI), and A/B fault testing."),
        ("Core Engineering Stack: ", "Go (Golang) for high-concurrency orchestrator routines, PostgreSQL 16 and MinIO for metadata and S3 object storage, and Next.js 14 / TypeScript for operator UI.")
    ]
    for bold_text, normal_text in techs:
        tp = doc.add_paragraph(style='List Bullet')
        tp.paragraph_format.left_indent = Inches(0.4)
        tp.paragraph_format.space_after = Pt(3)
        r_b = tp.add_run(bold_text)
        r_b.font.bold = True
        tp.add_run(normal_text)

    # 7. Work Plan
    add_section_header("7", "Work Plan")
    plan = [
        ("Phase 1: System Architecture & Security Modeling", "Formalizing edge state machine transitions, designing Canary slicing algorithms, and implementing the ECDSA P-256 cryptographic verification engine."),
        ("Phase 2: Orchestration & Edge Agent Engineering", "Developing the core deployment orchestrator, edge agent daemon, and resilient MQTT 5.0 telemetry communication channels."),
        ("Phase 3: Automated Testing & Fault Injection Framework", "Developing the multi-scenario automated test runner and fault-injection scripts for controlled simulation benchmarking."),
        ("Phase 4: Simulation Benchmarking & Statistical Evaluation", "Conducting multi-scenario experiments (N = 290 trials), running Mann-Whitney U hypothesis tests, and evaluating fleet survival rates."),
        ("Phase 5: Performance Verification & Thesis Finalization", "Analyzing edge computational overhead, documenting architectural findings, compiling performance charts, and defending the thesis.")
    ]
    
    plan_tbl = doc.add_table(rows=len(plan)+1, cols=2)
    plan_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = plan_tbl.rows[0].cells
    hdr_cells[0].text = "Project Milestone"
    hdr_cells[1].text = "Key Activities and Deliverables"
    hdr_cells[0].width = Inches(2.2)
    hdr_cells[1].width = Inches(4.3)
    for c in hdr_cells:
        set_cell_background(c, "1B365D")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for idx, (m, desc) in enumerate(plan, 1):
        row = plan_tbl.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        c0.text = m
        c1.text = desc
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        c0.paragraphs[0].runs[0].font.bold = True
        if idx % 2 == 1:
            set_cell_background(c0, "F5F7FA")
            set_cell_background(c1, "F5F7FA")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 8. Expected Benefits
    add_section_header("8", "Expected Benefits")
    bens = [
        ("Production Downtime Mitigation: ", "Mitigates manufacturing line disruptions caused by faulty software releases by restricting the failure blast radius to an initial canary group, significantly improving fleet survival rates and enabling automated rollback recovery compared to conventional deployments."),
        ("Zero-Trust Edge Cryptographic Integrity: ", "Provides a framework for evaluating cryptographic verification overhead on simulated edge nodes, with future validation required on representative industrial robot controllers."),
        ("Standardized Resilient Architecture: ", "Provides a robust, reproducible architectural blueprint for industrial automation engineers to deploy firmware updates safely across mission-critical, heterogeneous robot fleets.")
    ]
    for idx, (bold_text, normal_text) in enumerate(bens, 1):
        bp = doc.add_paragraph()
        bp.paragraph_format.left_indent = Inches(0.25)
        bp.paragraph_format.space_after = Pt(3)
        r_num = bp.add_run(f"8.{idx} ")
        r_num.font.bold = True
        r_b = bp.add_run(bold_text)
        r_b.font.bold = True
        bp.add_run(normal_text)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    try:
        doc.save(output_path)
        print(f"[SUCCESS] Saved updated Word document to: {output_path}")
    except PermissionError:
        print(f"[ERROR] Could not save to {output_path}. Please ensure the file is closed in Word.")

if __name__ == "__main__":
    out = os.path.abspath("docs/academic/PROJECT_PROPOSAL.docx")
    create_proposal_docx(out)

